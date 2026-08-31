"""文档解析适配器：pdfplumber / python-docx / 直读。

**不做 OCR**（spec §3.3）：扫描版 PDF（纯图片、无文字层）经 pdfplumber 解析后
文本为空或极短，此时判定为解析失败并给出提示，由 IndexingService 置
`status=failed` + `error_msg`，管理端可见并可重新上传。

所有解析都是同步阻塞调用，由调用方经 `run_in_threadpool` 卸载（ADR-0002）。
"""

from pathlib import Path
from typing import Protocol, runtime_checkable

from app.domain.knowledge.errors import EmptyDocumentError
from app.domain.knowledge.status import SOURCE_DOCX, SOURCE_MD, SOURCE_PDF, SOURCE_TXT
from app.infrastructure.ports.document import ParsedDocument

NO_TEXT_HINT = "未提取到文本，可能是扫描版 PDF，请先做 OCR"

# spec §8.2「文本为空或极短」的判定阈值。**仅对 PDF 生效**：该判据的目的是识别
# 无文字层的扫描版，对 txt / md / docx 而言文件正文就是全部内容，用长度阈值
# 会把合法的小文档误杀（例如只有百来字的单页讲义）。
MIN_PDF_TEXT_CHARS = 10

_ENCODING = "utf-8"


@runtime_checkable
class _FormatParser(Protocol):
    """适配器内部的逐格式接缝：只认文件本身，不认来源类型。"""

    def parse(self, path: Path) -> ParsedDocument: ...


class PlainTextParser:
    """Markdown 与 TXT 直接读文本，无需第三方库。"""

    def parse(self, path: Path) -> ParsedDocument:
        text = path.read_text(encoding=_ENCODING, errors="replace")
        return ParsedDocument(text=_require_non_empty(text), meta={"source_type": _kind(path)})


class PdfParser:
    """pdfplumber 逐页抽取文字层。"""

    def parse(self, path: Path) -> ParsedDocument:
        import pdfplumber  # 局部导入：未安装时仅 PDF 解析不可用，不影响其余类型

        pages: list[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        text = "\n".join(pages)
        _require_text_layer(text)
        return ParsedDocument(
            text=text.strip(), meta={"source_type": SOURCE_PDF, "pages": len(pages)}
        )


class DocxParser:
    """python-docx 逐段落抽取；表格内容一并纳入，否则讲义会整段丢失。"""

    def parse(self, path: Path) -> ParsedDocument:
        import docx  # 局部导入，理由同 PdfParser

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)

        text = "\n".join(parts)
        return ParsedDocument(text=_require_non_empty(text), meta={"source_type": SOURCE_DOCX})


_PARSERS: dict[str, _FormatParser] = {
    SOURCE_TXT: PlainTextParser(),
    SOURCE_MD: PlainTextParser(),
    SOURCE_PDF: PdfParser(),
    SOURCE_DOCX: DocxParser(),
}


def parse_document(path: Path, source_type: str) -> ParsedDocument:
    """按来源类型解析文档；无文字层或文件缺失抛领域/内置异常。"""
    try:
        parser = _PARSERS[source_type]
    except KeyError as exc:
        raise ValueError(f"不支持的来源类型：{source_type}") from exc

    if not path.exists():
        raise FileNotFoundError(f"源文件不存在：{path}")

    return parser.parse(path)


class MultiFormatDocumentParser:
    """`DocumentParser` 端口的正式实现：按 source_type 派发到逐格式解析器。

    与 `parse_document` 函数行为等价（本类就是它的类化外壳）；服务层经
    `runtime.get_document_parser()` 注入本类，不再直连适配器模块
    （spec §4.2 硬约束 2 / 健康检查 H-3）。
    """

    def parse(self, path: Path, source_type: str) -> ParsedDocument:
        return parse_document(path, source_type)


def _visible_chars(text: str) -> int:
    return len("".join(ch for ch in text if not ch.isspace()))


def _require_non_empty(text: str) -> str:
    if _visible_chars(text) == 0:
        raise EmptyDocumentError(NO_TEXT_HINT)
    return text.strip()


def _require_text_layer(text: str) -> None:
    """扫描版 PDF 的判据：去掉空白后仍短于阈值即视为无文字层。"""
    if _visible_chars(text) < MIN_PDF_TEXT_CHARS:
        raise EmptyDocumentError(NO_TEXT_HINT)


def _kind(path: Path) -> str:
    return SOURCE_MD if path.suffix.lower() == ".md" else SOURCE_TXT
